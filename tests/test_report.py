"""Приёмочные проверки XLSX-экспорта; запускаются через unittest."""

import copy
import importlib.util
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from zipfile import ZipFile

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1] / "skills" / "proverka-tekstov"
SPEC = importlib.util.spec_from_file_location("build_report", ROOT / "scripts/build_report.py")
REPORT = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(REPORT)


class ReportAcceptanceTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.directory = Path(self.temp.name)
        self.data = json.loads((ROOT / "assets/report-example.json").read_text())
        self.output = self.directory / "отчёт.xlsx"

    def tearDown(self):
        self.temp.cleanup()

    def test_demo_routes_every_record_and_preserves_exact_quotes(self):
        REPORT.build_report(self.data, self.output)
        wb = load_workbook(self.output)
        self.assertEqual(wb.sheetnames, ["Сводка", "Ошибки", "Уточнения", "Предложения", "Покрытие", "Факты"])
        self.assertEqual([wb[name].max_row - 4 for name in ("Ошибки", "Уточнения", "Предложения")], [3, 1, 1])
        found = {}
        for name in ("Ошибки", "Уточнения", "Предложения"):
            for row in wb[name].iter_rows(min_row=5, values_only=True):
                self.assertEqual(row[11], "Нет")
                found[row[0]] = row[3]
        self.assertEqual(found, {issue["id"]: issue["исходный_текст"] for issue in self.data["замечания"]})
        self.assertEqual(wb["Сводка"]["B5"].value, 3)
        self.assertEqual(wb["Сводка"]["B7"].value, 1)
        self.assertEqual(wb["Сводка"]["B9"].value, 5)
        self.assertEqual(wb["Сводка"]["B10"].value, 1)
        self.assertEqual(wb["Сводка"]["B11"].value, 1)
        self.assertEqual(wb["Факты"]["E5"].number_format, "dd.mm.yyyy")
        for name in ("Ошибки", "Уточнения", "Предложения"):
            sheet = wb[name]
            self.assertEqual(sheet.freeze_panes, "C5")
            self.assertTrue(sheet.auto_filter.ref)
            self.assertEqual(len(sheet.data_validations.dataValidation), 1)
            validation = sheet.data_validations.dataValidation[0]
            self.assertEqual(validation.formula1, '"Да,Нет"')
            self.assertTrue(validation.showErrorMessage)
            self.assertIn("L5", validation.sqref)
        wb.close()

    def test_formula_like_material_is_written_as_literal_without_prefix(self):
        attacks = ['=HYPERLINK("https://example.invalid","открыть")', "+SUM(1,2)", "-2+3", "@SUM(1,2)", "  =1+1", "\t=1+1"]
        template = self.data["замечания"][0]
        self.data["замечания"] = []
        for number, attack in enumerate(attacks):
            issue = copy.deepcopy(template)
            issue["id"] = str(number)
            issue["исходный_текст"] = attack
            issue["исправление"] = attack
            issue["источник"] = attack
            self.data["замечания"].append(issue)
        self.data["название"] = "=1+1"
        self.data["источники"][0]["название"] = "=1+2"
        self.data["покрытие"][0]["причина"] = "=1+3"
        self.data["факты"][0]["источник"] = "=1+4"
        REPORT.build_report(self.data, self.output)
        wb = load_workbook(self.output, data_only=False)
        for row, attack in enumerate(attacks, 5):
            for column in (4, 8, 10):
                cell = wb["Ошибки"].cell(row, column)
                self.assertEqual(cell.value, attack)
                self.assertEqual(cell.data_type, "s")
        for sheet in wb:
            for row in sheet:
                for cell in row:
                    self.assertNotEqual(cell.data_type, "f")
        self.assertEqual(wb["Сводка"]["A2"].value, "=1+1")
        self.assertEqual(wb["Сводка"]["A18"].value, "=1+2")
        wb.close()
        with ZipFile(self.output) as archive:
            for name in archive.namelist():
                if name.startswith("xl/worksheets/sheet") and name.endswith(".xml"):
                    self.assertNotIn(b"<f>", archive.read(name))

    def test_bad_status_is_rejected_and_existing_output_survives(self):
        self.output.write_bytes(b"existing report")
        self.data["замечания"][0]["статус"] = "Исправить потом"
        with self.assertRaisesRegex(REPORT.ReportError, "недопустимое значение"):
            REPORT.build_report(self.data, self.output)
        self.assertEqual(self.output.read_bytes(), b"existing report")

    def test_unknown_field_and_duplicate_ids_do_not_disappear(self):
        self.data["замечания"][0]["авторский_комментарий"] = "Не потерять"
        with self.assertRaisesRegex(REPORT.ReportError, "неизвестные поля"):
            REPORT.validate_report(self.data)
        del self.data["замечания"][0]["авторский_комментарий"]
        self.data["замечания"][1]["id"] = self.data["замечания"][0]["id"]
        with self.assertRaisesRegex(REPORT.ReportError, "повторный идентификатор"):
            REPORT.validate_report(self.data)

    def test_max_length_is_preserved_and_overflow_is_rejected(self):
        self.data["замечания"][0]["исходный_текст"] = "А" * 32767
        REPORT.build_report(self.data, self.output)
        wb = load_workbook(self.output)
        self.assertEqual(wb["Ошибки"]["D5"].value, "А" * 32767)
        wb.close()
        for long_text in ("А" * 32768, "😀" * 16384):
            self.data["замечания"][0]["исходный_текст"] = long_text
            with self.assertRaisesRegex(REPORT.ReportError, "текст превышает лимит Excel"):
                REPORT.validate_report(self.data)

    def test_invalid_xml_character_fails_without_silent_removal(self):
        self.data["замечания"][0]["исходный_текст"] = "Текст\x00ещё"
        with self.assertRaisesRegex(REPORT.ReportError, "недопустимый для XLSX символ"):
            REPORT.validate_report(self.data)

    def test_partial_coverage_requires_reason_and_facts_require_evidence(self):
        self.data["покрытие"][1]["причина"] = ""
        with self.assertRaisesRegex(REPORT.ReportError, "укажите, почему"):
            REPORT.validate_report(self.data)
        self.data["покрытие"][1]["причина"] = "Неразборчиво"
        self.data["факты"][0]["источник"] = ""
        with self.assertRaisesRegex(REPORT.ReportError, "необходим источник"):
            REPORT.validate_report(self.data)

    def test_empty_results_are_valid_and_do_not_imply_full_coverage(self):
        self.data = {"название": "Пустой отчёт", "источники": [], "замечания": [], "покрытие": [], "факты": []}
        REPORT.build_report(self.data, self.output)
        wb = load_workbook(self.output)
        self.assertEqual(wb["Сводка"]["B5"].value, 0)
        self.assertEqual(wb["Сводка"]["B10"].value, 0)
        self.assertEqual(wb["Ошибки"].max_row, 4)
        self.assertIn("покрытие не описано", wb["Покрытие"]["A2"].value)
        wb.close()

    def test_cli_and_duplicate_json_keys(self):
        input_file = self.directory / "данные.json"
        input_file.write_text(json.dumps(self.data, ensure_ascii=False), encoding="utf-8")
        result = subprocess.run([sys.executable, str(ROOT / "scripts/build_report.py"), str(input_file), "--output", str(self.output)], capture_output=True, text=True)
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertIn("Отчёт сохранён", result.stdout)
        input_file.write_text('{"название":"а","название":"б"}', encoding="utf-8")
        result = subprocess.run([sys.executable, str(ROOT / "scripts/build_report.py"), str(input_file), "--output", str(self.output)], capture_output=True, text=True)
        self.assertEqual(result.returncode, 2)
        self.assertIn("повторный ключ", result.stderr)


if __name__ == "__main__":
    unittest.main()

"""Проверки адресов, пропусков и неизменности входных файлов."""
import hashlib
import importlib.util
import json
import subprocess
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

SCRIPT = Path(__file__).resolve().parents[1] / "skills" / "proverka-tekstov" / "scripts" / "extract_text.py"
spec = importlib.util.spec_from_file_location('extract_text', SCRIPT)
module = importlib.util.module_from_spec(spec)
spec.loader.exec_module(module)


class ExtractorTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)

    def tearDown(self):
        self.tmp.cleanup()

    def run_extractor(self, path):
        before = hashlib.sha256(path.read_bytes()).hexdigest()
        data = module.Extraction(path).run()
        self.assertEqual(before, hashlib.sha256(path.read_bytes()).hexdigest())
        self.assertEqual(data['источники'][0]['статус_извлечения'], 'извлечено частично')
        self.assertTrue(all(c['визуальная_проверка'] == 'не выполнена' for c in data['покрытие']))
        self.assertTrue(all({'источник', 'место', 'текст', 'способ', 'служебный'} <= set(f)
                            for f in data['фрагменты']))
        json.dumps(data, ensure_ascii=False, allow_nan=False)
        return data

    def test_xlsx_far_cell_hidden_sheet_formula_and_cache(self):
        from openpyxl import Workbook
        w = Workbook()
        ws = w.active
        ws.title = 'Заказы'
        ws['A1'] = 'Заголовок'
        ws['C10500'] = 'Последняя строка'
        ws['B2'] = '=2+3'
        ws['B3'] = '=2+4'
        ws.row_dimensions[10500].hidden = True
        ws.column_dimensions.group('C', 'E', hidden=True)
        hidden = w.create_sheet('Архив')
        hidden.sheet_state = 'hidden'
        hidden['Z1001'] = 'Скрытая запись'
        path = self.root / 'book.xlsx'
        w.save(path)
        with zipfile.ZipFile(path) as z:
            files = {name: z.read(name) for name in z.namelist()}
        namespace = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
        root = ET.fromstring(files['xl/worksheets/sheet1.xml'])
        cached = root.find('.//s:c[@r="B2"]/s:v', namespace)
        cached.text = '5'
        files['xl/worksheets/sheet1.xml'] = ET.tostring(root)
        with zipfile.ZipFile(path, 'w') as z:
            for name, payload in files.items():
                z.writestr(name, payload)
        data = self.run_extractor(path)
        texts = {f['текст']: f for f in data['фрагменты']}
        self.assertIn('Последняя строка', texts)
        self.assertIn('C10500', texts['Последняя строка']['место'])
        self.assertTrue(texts['Последняя строка']['скрытая_строка'])
        self.assertTrue(texts['Последняя строка']['скрытый_столбец'])
        self.assertTrue(texts['Скрытая запись']['служебный'])
        self.assertEqual(texts['5']['способ'], 'кэш формулы XLSX')
        self.assertEqual(len([l for l in data['ограничения'] if l['вид'] == 'значение не пересчитано']), 2)
        self.assertEqual(data['покрытие'][0]['непустых_ячеек'], 4)

    def test_docx_tables_hyperlinks_headers_and_images(self):
        from docx import Document
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from PIL import Image
        d = Document()
        d.add_heading('Заголовок раздела', 1)
        paragraph = d.add_paragraph('Перед ссылкой ')
        link = OxmlElement('w:hyperlink')
        run = OxmlElement('w:r')
        text = OxmlElement('w:t')
        text.text = 'Текст ссылки'
        run.append(text)
        link.append(run)
        paragraph._p.append(link)
        table = d.add_table(rows=1, cols=2)
        table.cell(0, 0).text = 'Ячейка слева'
        nested = table.cell(0, 1).add_table(rows=1, cols=1)
        nested.cell(0, 0).text = 'Вложенная ячейка'
        d.sections[0].header.paragraphs[0].text = 'Верх страницы'
        d.sections[0].footer.paragraphs[0].text = 'Низ страницы'
        png = self.root / 'picture.png'
        Image.new('RGB', (4, 4), 'white').save(png)
        d.add_picture(str(png))
        path = self.root / 'document.docx'
        d.save(path)
        data = self.run_extractor(path)
        texts = {f['текст']: f for f in data['фрагменты']}
        self.assertIn('Перед ссылкой Текст ссылки', texts)
        self.assertEqual(texts['Заголовок раздела']['стиль_абзаца'], 'Heading1')
        self.assertIn('таблица 2', texts['Вложенная ячейка']['место'])
        self.assertIn('Верх страницы', texts)
        self.assertIn('Низ страницы', texts)
        self.assertTrue(any('рисунок или схема' in l['место'] for l in data['ограничения']))

    def test_pptx_hidden_slide_group_table_notes_and_image(self):
        from pptx import Presentation
        from pptx.util import Inches
        from PIL import Image
        p = Presentation()
        slide = p.slides.add_slide(p.slide_layouts[6])
        slide._element.set('show', '0')
        group = slide.shapes.add_group_shape()
        box = group.shapes.add_textbox(0, 0, Inches(2), Inches(1))
        box.text_frame.text = 'Внутри группы'
        table = slide.shapes.add_table(1, 1, 0, Inches(2), Inches(3), Inches(1)).table
        table.cell(0, 0).text = 'Табличная запись'
        slide.notes_slide.notes_text_frame.text = 'Заметка докладчика'
        png = self.root / 'picture.png'
        Image.new('RGB', (4, 4), 'white').save(png)
        slide.shapes.add_picture(str(png), 0, Inches(3))
        path = self.root / 'deck.pptx'
        p.save(path)
        data = self.run_extractor(path)
        texts = {f['текст']: f for f in data['фрагменты']}
        self.assertTrue(texts['Внутри группы']['скрытый_слайд'])
        self.assertEqual(texts['Внутри группы']['место'].count('объект'), 2)
        self.assertTrue(texts['Табличная запись']['служебный'])
        self.assertTrue(texts['Заметка докладчика']['служебный'])
        self.assertTrue(any('на изображении' in l['описание'] for l in data['ограничения']))

    def test_docx_preserves_no_break_hyphen_element(self):
        from docx import Document
        from docx.oxml import OxmlElement
        d = Document()
        paragraph = d.add_paragraph()
        run = paragraph.add_run('бизнес')
        run._r.append(OxmlElement('w:noBreakHyphen'))
        run.add_text('процесс')
        path = self.root / 'nonbreaking.docx'
        d.save(path)
        data = self.run_extractor(path)
        self.assertEqual(data['фрагменты'][0]['текст'], 'бизнес\u2011процесс')

    def test_pdf_text_bbox_and_scanned_page(self):
        import fitz
        from PIL import Image
        d = fitz.open()
        p1 = d.new_page()
        p1.insert_text((50, 50), 'Known text')
        png = self.root / 'picture.png'
        Image.new('RGB', (100, 100), 'white').save(png)
        p2 = d.new_page()
        p2.insert_image(fitz.Rect(0, 0, 100, 100), filename=str(png))
        path = self.root / 'document.pdf'
        d.save(path)
        d.close()
        data = self.run_extractor(path)
        fragment = next(f for f in data['фрагменты'] if 'Known text' in f['текст'])
        self.assertEqual(len(fragment['координаты']['прямоугольник']), 4)
        self.assertEqual(data['покрытие'][1]['изображений'], 1)
        self.assertTrue(any(l['место'] == 'Страница 2' and 'текст не найден' in l['описание']
                            for l in data['ограничения']))

    def test_png_and_jpeg_without_russian_ocr_are_honest(self):
        from PIL import Image
        from unittest.mock import patch
        for suffix in ('.png', '.jpeg'):
            path = self.root / ('image' + suffix)
            Image.new('RGB', (50, 50), 'white').save(path)
            result = subprocess.CompletedProcess([], 0, 'List of available languages (1):\neng\n', '')
            with patch.object(module.shutil, 'which', return_value='/bin/tesseract'):
                with patch.object(module.subprocess, 'run', return_value=result) as mocked:
                    data = self.run_extractor(path)
                    self.assertEqual(mocked.call_count, 1)
            self.assertFalse(data['фрагменты'])
            self.assertTrue(any('OCR не запускался' in l['описание'] for l in data['ограничения']))

    def test_cli_rejects_overwriting_input_and_writes_read_failure(self):
        path = self.root / 'broken.pdf'
        path.write_bytes(b'not pdf')
        before = path.read_bytes()
        result = subprocess.run([sys.executable, str(SCRIPT), str(path), '--output', str(path)], capture_output=True)
        self.assertEqual(result.returncode, 2)
        self.assertEqual(path.read_bytes(), before)
        output = self.root / 'result.json'
        result = subprocess.run([sys.executable, str(SCRIPT), str(path), '--output', str(output)], capture_output=True)
        self.assertEqual(result.returncode, 1)
        self.assertEqual(json.loads(output.read_text())['источники'][0]['статус_извлечения'], 'ошибка извлечения')

    def test_ocr_keeps_bbox_and_requires_visual_confirmation(self):
        from PIL import Image
        from unittest.mock import patch
        path = self.root / 'image.png'
        Image.new('RGB', (50, 50), 'white').save(path)
        languages = subprocess.CompletedProcess([], 0, 'List of available languages (2):\nrus\neng\n', '')
        tsv = ('level\tpage_num\tblock_num\tpar_num\tline_num\tword_num\tleft\ttop\twidth\theight\tconf\ttext\n'
               '5\t1\t1\t1\t1\t1\t10\t10\t20\t10\t90\tТекст\n'
               '5\t1\t1\t1\t1\t2\t35\t10\t10\t10\t75\tтут\n')
        recognized = subprocess.CompletedProcess([], 0, tsv, '')
        with patch.object(module.shutil, 'which', return_value='/bin/tesseract'):
            with patch.object(module.subprocess, 'run', side_effect=[languages, recognized]):
                data = self.run_extractor(path)
        self.assertEqual(data['фрагменты'][0]['текст'], 'Текст тут')
        self.assertEqual(data['фрагменты'][0]['координаты']['прямоугольник'], [10, 10, 45, 20])
        self.assertEqual(data['фрагменты'][0]['уверенность_OCR'], 75)
        self.assertIn('требует сверки', data['фрагменты'][0]['способ'])

    def test_png_alpha_is_composited_before_ocr(self):
        from PIL import Image
        from unittest.mock import patch
        path = self.root / 'transparent.png'
        im = Image.new('RGBA', (10, 10), (0, 0, 0, 0))
        im.putpixel((3, 3), (0, 0, 0, 255))
        im.save(path)
        ocr_calls = []

        def tesseract(args, **kwargs):
            if '--list-langs' in args:
                return subprocess.CompletedProcess(args, 0, 'List of available languages (1):\nrus\n', '')
            with Image.open(args[1]) as prepared:
                self.assertEqual(prepared.mode, 'RGB')
                self.assertEqual(prepared.getpixel((0, 0)), (255, 255, 255))
                self.assertEqual(prepared.getpixel((3, 3)), (0, 0, 0))
            ocr_calls.append(args)
            return subprocess.CompletedProcess(args, 0, 'level\ttext\n', '')

        with patch.object(module.shutil, 'which', return_value='/bin/tesseract'):
            with patch.object(module.subprocess, 'run', side_effect=tesseract):
                data = self.run_extractor(path)
        self.assertEqual(len(ocr_calls), 1)
        self.assertIn('белый', data['источники'][0]['фон_для_OCR'])
        self.assertTrue(any('может скрыть светлый текст' in l['описание'] for l in data['ограничения']))


if __name__ == '__main__':
    unittest.main(verbosity=2)

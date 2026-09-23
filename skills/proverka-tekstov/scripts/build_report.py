#!/usr/bin/env python3
"""Собрать отчёт о проверке текстов в XLSX из JSON. Требуется openpyxl 3.1+."""

from __future__ import annotations

import argparse
from collections import Counter
from datetime import date
import json
import math
import os
from pathlib import Path
import sys
import tempfile
from typing import Any


STATUSES = (
    "Ошибка подтверждена",
    "Нужно уточнение",
    "Факт не проверен",
    "Редакторское предложение",
)
GROUPS = (
    "Грамматические ошибки",
    "Речевые ошибки",
    "Логические и фактические ошибки",
)
PRIORITIES = ("Высокий", "Обычный")
COVERAGE_STATUSES = ("Проверено полностью", "Проверено частично", "Не проверено")
FACT_STATUSES = ("Подтверждено", "Опровергнуто", "Факт не проверен", "Нужно уточнение")
TOP_FIELDS = ("название", "источники", "замечания", "покрытие", "факты")
ISSUE_FIELDS = (
    "id", "файл", "место", "исходный_текст", "группа", "тип", "статус",
    "исправление", "пояснение", "источник", "приоритет",
)
SOURCE_FIELDS = ("название", "область")
COVERAGE_FIELDS = ("файл", "область", "статус", "причина")
FACT_FIELDS = ("утверждение", "место", "статус", "источник", "дата_проверки")
MAX_CELL_UNITS = 32767
MAX_ROWS = 1048576
HEADER_ROW = 4
FIRST_DATA_ROW = HEADER_ROW + 1


class ReportError(ValueError):
    """Ошибка входных данных с объяснением на русском языке."""


def _exact_fields(value: Any, expected: tuple[str, ...], path: str) -> None:
    if not isinstance(value, dict):
        raise ReportError(f"{path}: требуется объект JSON.")
    missing = set(expected) - value.keys()
    extra = value.keys() - set(expected)
    if missing:
        raise ReportError(f"{path}: отсутствуют обязательные поля: {', '.join(sorted(missing))}.")
    if extra:
        raise ReportError(f"{path}: неизвестные поля: {', '.join(sorted(extra))}. Данные не будут отброшены молча.")


def _text(value: Any, path: str, *, allow_empty: bool = False) -> None:
    if not isinstance(value, str):
        raise ReportError(f"{path}: требуется строка.")
    if not allow_empty and not value.strip():
        raise ReportError(f"{path}: строка не должна быть пустой.")
    # Excel ограничивает текст ячейки; суррогатные пары считаем консервативно.
    for char in value:
        code = ord(char)
        if (code < 32 and char not in "\t\n\r") or 0xD800 <= code <= 0xDFFF or code in (0xFFFE, 0xFFFF):
            raise ReportError(f"{path}: недопустимый для XLSX символ U+{code:04X}. Исправьте исходный JSON; текст не удалён автоматически.")
    units = len(value.encode("utf-16-le")) // 2
    if units > MAX_CELL_UNITS:
        raise ReportError(
            f"{path}: текст превышает лимит Excel {MAX_CELL_UNITS} кодовых единиц UTF-16 "
            f"({units}). Разделите запись на несколько замечаний с отдельными id; текст не обрезан."
        )


def _choice(value: str, choices: tuple[str, ...], path: str) -> None:
    if value not in choices:
        raise ReportError(f"{path}: недопустимое значение «{value}». Допустимо: {', '.join(choices)}.")


def _rows(data: dict[str, Any], key: str, fields: tuple[str, ...], empty: set[str]) -> None:
    rows = data[key]
    if not isinstance(rows, list):
        raise ReportError(f"{key}: требуется массив.")
    max_items = MAX_ROWS - (17 if key == "источники" else HEADER_ROW)
    if len(rows) > max_items:
        raise ReportError(f"{key}: слишком много записей для листа Excel. Максимум: {max_items}.")
    for index, row in enumerate(rows, 1):
        path = f"{key}[{index}]"
        _exact_fields(row, fields, path)
        for field in fields:
            _text(row[field], f"{path}.{field}", allow_empty=field in empty)


def validate_report(data: Any) -> dict[str, Any]:
    """Проверить весь документ до создания файла, без потери неизвестных данных."""
    _exact_fields(data, TOP_FIELDS, "Отчёт")
    _text(data["название"], "название")
    _rows(data, "источники", SOURCE_FIELDS, set())
    _rows(data, "замечания", ISSUE_FIELDS, {"исправление", "источник"})
    _rows(data, "покрытие", COVERAGE_FIELDS, {"причина"})
    _rows(data, "факты", FACT_FIELDS, {"источник", "дата_проверки"})
    seen = set()
    for index, issue in enumerate(data["замечания"], 1):
        path = f"замечания[{index}]"
        if issue["id"] in seen:
            raise ReportError(f"{path}.id: повторный идентификатор «{issue['id']}».")
        seen.add(issue["id"])
        _choice(issue["статус"], STATUSES, f"{path}.статус")
        _choice(issue["группа"], GROUPS, f"{path}.группа")
        _choice(issue["приоритет"], PRIORITIES, f"{path}.приоритет")
    for index, coverage in enumerate(data["покрытие"], 1):
        _choice(coverage["статус"], COVERAGE_STATUSES, f"покрытие[{index}].статус")
        if coverage["статус"] != "Проверено полностью" and not coverage["причина"].strip():
            raise ReportError(f"покрытие[{index}].причина: укажите, почему область проверена не полностью.")
    for index, fact in enumerate(data["факты"], 1):
        path = f"факты[{index}]"
        _choice(fact["статус"], FACT_STATUSES, f"{path}.статус")
        checked = fact["дата_проверки"]
        if checked:
            try:
                if len(checked) != 10 or date.fromisoformat(checked).isoformat() != checked:
                    raise ValueError()
            except ValueError:
                raise ReportError(f"{path}.дата_проверки: используйте существующую дату ГГГГ-ММ-ДД или пустую строку.") from None
        if fact["статус"] in ("Подтверждено", "Опровергнуто") and not fact["источник"].strip():
            raise ReportError(f"{path}.источник: для проверенного факта необходим источник или расчёт.")
    return data


def _reject_duplicate_keys(pairs: list[tuple[str, Any]]) -> dict[str, Any]:
    result: dict[str, Any] = {}
    for key, value in pairs:
        if key in result:
            raise ReportError(f"JSON содержит повторный ключ «{key}». Удалите неоднозначность.")
        result[key] = value
    return result


def read_report(path: Path) -> dict[str, Any]:
    try:
        text = path.read_text(encoding="utf-8-sig")
    except (OSError, UnicodeError) as exc:
        raise ReportError(f"Не удалось прочитать JSON в кодировке UTF-8: {exc}") from exc
    try:
        data = json.loads(text, object_pairs_hook=_reject_duplicate_keys)
    except json.JSONDecodeError as exc:
        raise ReportError(f"Некорректный JSON: строка {exc.lineno}, столбец {exc.colno}.") from exc
    return validate_report(data)


def _write_text(cell: Any, value: str) -> None:
    # Явный строковый тип сохраняет =,+,-,@ без апострофа и выполнения формул.
    cell.value = value
    cell.data_type = "s"
    cell.number_format = "@"


def build_report(data: Any, output: Path) -> None:
    """Создать XLSX атомарно: существующий результат не портится при ошибке."""
    validate_report(data)
    output = Path(output)
    if output.suffix.lower() != ".xlsx":
        raise ReportError("Выходной файл должен иметь расширение .xlsx.")
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.worksheet.datavalidation import DataValidation
        from openpyxl.utils import get_column_letter
    except ImportError as exc:
        raise ReportError("Не установлена зависимость openpyxl. Установите зависимости из requirements.txt плагина.") from exc

    wb = Workbook()
    wb.remove(wb.active)
    wb.properties.title = data["название"]
    wb.properties.subject = "Результаты проверки текстов"
    wb.properties.creator = "Проверка текстов"
    navy, blue, pale, gray = "17324D", "245C80", "F2F6FA", "526171"
    bottom = Border(bottom=Side(style="hair", color="DCE3EA"))

    def base_sheet(name: str, subtitle: str, widths: list[float]):
        ws = wb.create_sheet(name)
        ws.sheet_view.showGridLines = False
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.sheet_properties.outlinePr.summaryRight = False
        ws.sheet_properties.tabColor = blue
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A3
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.freeze_panes = "C5" if name in ("Ошибки", "Уточнения", "Предложения") else "A5"
        for index, width in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(index)].width = width
        _write_text(ws.cell(1, 1), name)
        ws.cell(1, 1).font = Font(name="Aptos", size=20, bold=True, color=navy)
        ws.row_dimensions[1].height = 31
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(widths))
        _write_text(ws.cell(2, 1), subtitle)
        ws.cell(2, 1).font = Font(name="Aptos", size=10, color=gray)
        ws.cell(2, 1).alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[2].height = 32
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(widths))
        return ws

    def table(ws, headers: list[str], rows: list[list[Any]], widths: list[float], header_row=HEADER_ROW):
        for column, header in enumerate(headers, 1):
            cell = ws.cell(header_row, column)
            _write_text(cell, header)
            cell.fill = PatternFill("solid", fgColor=navy)
            cell.font = Font(name="Aptos", size=10, bold=True, color="FFFFFF")
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.row_dimensions[header_row].height = 32
        for row_number, values in enumerate(rows, header_row + 1):
            lines = 1
            for column, value in enumerate(values, 1):
                cell = ws.cell(row_number, column)
                if isinstance(value, str):
                    _write_text(cell, value)
                    line_count = sum(max(1, math.ceil(len(part) / max(1, int(widths[column - 1] - 3)))) for part in value.split("\n"))
                    lines = max(lines, line_count)
                else:
                    cell.value = value
                cell.font = Font(name="Aptos", size=10, color=navy)
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                cell.border = bottom
                if row_number % 2:
                    cell.fill = PatternFill("solid", fgColor=pale)
            ws.row_dimensions[row_number].height = min(409, max(31, 14 * lines + 9))
        last_row = max(header_row, header_row + len(rows))
        ws.auto_filter.ref = f"A{header_row}:{get_column_letter(len(headers))}{last_row}"
        ws.print_title_rows = f"1:{header_row}"
        ws.print_options.horizontalCentered = False
        ws.print_area = f"A1:{get_column_letter(len(headers))}{max(last_row, header_row + 1)}"

    counts = Counter(item["статус"] for item in data["замечания"])
    coverage_counts = Counter(item["статус"] for item in data["покрытие"])
    summary = base_sheet("Сводка", data["название"], [46, 105])
    summary.freeze_panes = "A5"
    summary_rows = [
        ["Подтверждённые ошибки", counts[STATUSES[0]]],
        ["Нужно уточнение", counts[STATUSES[1]]],
        ["Факты без подтверждения", counts[STATUSES[2]]],
        ["Редакторские предложения", counts[STATUSES[3]]],
        ["Всего замечаний", len(data["замечания"])],
        ["Областей проверено полностью", coverage_counts[COVERAGE_STATUSES[0]]],
        ["Областей проверено частично", coverage_counts[COVERAGE_STATUSES[1]]],
        ["Областей не проверено", coverage_counts[COVERAGE_STATUSES[2]]],
        ["Учёт покрытия", "Область — строка листа «Покрытие». Количество областей не равно количеству страниц или файлов."],
        ["Отметки исправления", "На листах замечаний выберите «Да» или «Нет» в столбце «Исправлено». Сводка отражает проверку на момент формирования и не пересчитывается при отметках."],
        ["Полный текст", "Цитаты сохранены целиком. Если длинный текст не помещается по высоте строки, откройте ячейку в строке формул Excel."],
    ]
    table(summary, ["Показатель", "Значение"], summary_rows, [46, 105])
    table(summary, ["Материал", "Область проверки"], [[source[k] for k in SOURCE_FIELDS] for source in data["источники"]], [46, 105], header_row=17)
    summary.auto_filter.ref = f"A17:B{max(17, 17 + len(data['источники']))}"
    summary.print_title_rows = "1:4"

    issue_headers = ["ID", "Файл", "Место", "Исходный текст", "Группа", "Тип", "Статус", "Исправление", "Пояснение", "Источник / расчёт", "Приоритет", "Исправлено"]
    issue_widths = [12, 25, 32, 48, 28, 24, 27, 48, 56, 48, 16, 15]
    partitions = [
        ("Ошибки", {STATUSES[0]}, "Подтверждённые ошибки. Исправления сохраняют смысл исходного текста."),
        ("Уточнения", {STATUSES[1], STATUSES[2]}, "Вопросы автору и факты без достаточного подтверждения. Статус сохранён у каждого замечания."),
        ("Предложения", {STATUSES[3]}, "Необязательные редакторские изменения. Исходный вариант допустим."),
    ]
    for name, selected, subtitle in partitions:
        rows = [[issue[key] for key in ISSUE_FIELDS] + ["Нет"] for issue in data["замечания"] if issue["статус"] in selected]
        ws = base_sheet(name, subtitle, issue_widths)
        table(ws, issue_headers, rows, issue_widths)
        validation = DataValidation(type="list", formula1='"Да,Нет"', allow_blank=False)
        validation.errorTitle = "Недопустимое значение"
        validation.error = "Выберите «Да» или «Нет» из списка."
        validation.promptTitle = "Исправлено"
        validation.prompt = "Отметьте, внесено ли исправление."
        validation.showErrorMessage = True
        validation.errorStyle = "stop"
        validation.showInputMessage = True
        validation.showDropDown = False
        ws.add_data_validation(validation)
        validation.add(f"L{FIRST_DATA_ROW}:L{max(FIRST_DATA_ROW, HEADER_ROW + len(rows))}")
        for row_number in range(FIRST_DATA_ROW, FIRST_DATA_ROW + len(rows)):
            ws.cell(row_number, 12).fill = PatternFill("solid", fgColor="EAF3FF")
            if ws.cell(row_number, 11).value == "Высокий":
                ws.cell(row_number, 11).font = Font(name="Aptos", size=10, bold=True, color="A12622")

    coverage_widths = [28, 60, 28, 95]
    ws = base_sheet("Покрытие", "Области проверки и причины пропусков. Пустой лист означает, что покрытие не описано.", coverage_widths)
    table(ws, ["Файл", "Область", "Статус", "Причина / ограничение"], [[item[key] for key in COVERAGE_FIELDS] for item in data["покрытие"]], coverage_widths)
    fact_widths = [65, 44, 28, 70, 20]
    ws = base_sheet("Факты", "Основания проверки утверждений. Отсутствие подтверждения само по себе не означает ошибку.", fact_widths)
    table(ws, ["Утверждение", "Место", "Статус", "Источник / расчёт", "Дата проверки"], [[item[key] for key in FACT_FIELDS] for item in data["факты"]], fact_widths)
    # Даты проверки — отдельное структурированное поле, а не цитаты материала.
    for row_number, item in enumerate(data["факты"], FIRST_DATA_ROW):
        if item["дата_проверки"]:
            cell = ws.cell(row_number, 5)
            cell.value = date.fromisoformat(item["дата_проверки"])
            cell.number_format = "dd.mm.yyyy"

    temporary: Path | None = None
    try:
        output.parent.mkdir(parents=True, exist_ok=True)
        with tempfile.NamedTemporaryFile(prefix=".report-", suffix=".xlsx", dir=output.parent, delete=False) as handle:
            temporary = Path(handle.name)
        wb.save(temporary)
        os.replace(temporary, output)
    except OSError as exc:
        raise ReportError(f"Не удалось сохранить отчёт: {exc}") from exc
    finally:
        wb.close()
        if temporary is not None:
            temporary.unlink(missing_ok=True)


class RussianArgumentParser(argparse.ArgumentParser):
    def format_usage(self) -> str:
        return super().format_usage().replace("usage: ", "Использование: ", 1)

    def format_help(self) -> str:
        return super().format_help().replace("usage: ", "Использование: ", 1)

    def error(self, message: str) -> None:
        self.print_usage(sys.stderr)
        self.exit(2, "Ошибка параметров. Укажите входной JSON и --output с путём к XLSX. Справка: --help.\n")


def main(argv: list[str] | None = None) -> int:
    parser = RussianArgumentParser(
        description="Собрать отчёт о проверке текстов из JSON в XLSX.",
        usage="%(prog)s ВХОДНОЙ_JSON --output ОТЧЁТ.xlsx",
        add_help=False,
    )
    parser._positionals.title = "Входные данные"
    parser._optionals.title = "Параметры"
    parser.add_argument("input", metavar="ВХОДНОЙ_JSON", type=Path, help="Путь к JSON в кодировке UTF-8.")
    parser.add_argument("--output", required=True, metavar="ОТЧЁТ.xlsx", type=Path, help="Путь для сохранения XLSX.")
    parser.add_argument("-h", "--help", action="help", help="Показать справку и завершить работу.")
    args = parser.parse_args(argv)
    try:
        if args.input.resolve() == args.output.resolve():
            raise ReportError("Входной JSON и выходной XLSX должны иметь разные пути.")
        data = read_report(args.input)
        build_report(data, args.output)
    except ReportError as exc:
        print(f"Ошибка: {exc}", file=sys.stderr)
        return 2
    print(f"Отчёт сохранён: {args.output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

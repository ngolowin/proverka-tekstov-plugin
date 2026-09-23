#!/usr/bin/env python3
"""Извлекает доступный текст и адреса; не выполняет редакторскую проверку.

Использование: python extract_text.py INPUT --output OUTPUT
Поддержка: PDF, DOCX, PPTX, XLSX, JPEG, PNG. Сетевых запросов нет.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import io
import json
import shutil
import subprocess
import sys
import tempfile
import zipfile
from datetime import date, datetime, time
from pathlib import Path
from typing import Any
from xml.etree import ElementTree as ET


VERSION = "1.0"
SUPPORTED = {".pdf", ".docx", ".pptx", ".xlsx", ".jpg", ".jpeg", ".png"}
W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


class Extraction:
    def __init__(self, path: Path):
        self.path = path
        self.data: dict[str, Any] = {
            "версия_схемы": VERSION,
            "назначение": "Извлечение текста; лингвистическая и фактическая проверка не выполнены",
            "источники": [{
                "идентификатор": "ист-1", "имя": path.name,
                "формат": path.suffix.lower().lstrip("."),
                "статус_извлечения": "не начато",
            }],
            "фрагменты": [], "покрытие": [], "ограничения": [],
        }

    def fragment(self, place: str, text: Any, method: str, service=False, **extra):
        if text is None or not str(text).strip():
            return
        self.data["фрагменты"].append({
            "идентификатор": f"фр-{len(self.data['фрагменты']) + 1}",
            "источник": "ист-1", "место": place, "текст": str(text),
            "способ": method, "служебный": bool(service), **extra,
        })

    def limitation(self, place: str, reason: str, kind="требуется визуальная проверка"):
        self.data["ограничения"].append({
            "источник": "ист-1", "место": place, "вид": kind, "описание": reason,
        })

    def coverage(self, place: str, **extra):
        self.data["покрытие"].append({
            "источник": "ист-1", "место": place,
            "статус_извлечения": "извлечено частично",
            "визуальная_проверка": "не выполнена", **extra,
        })

    def run(self):
        if not self.path.is_file():
            raise ValueError("Входной файл не найден или это каталог")
        suffix = self.path.suffix.lower()
        if suffix not in SUPPORTED:
            raise ValueError("Формат не поддерживается: " + suffix)
        digest = hashlib.sha256()
        with self.path.open("rb") as f:
            for chunk in iter(lambda: f.read(1024 * 1024), b""):
                digest.update(chunk)
        self.data["источники"][0]["sha256"] = digest.hexdigest()
        handler = {
            ".pdf": self.pdf, ".docx": self.docx, ".pptx": self.pptx,
            ".xlsx": self.xlsx, ".png": self.image, ".jpg": self.image, ".jpeg": self.image,
        }[suffix]
        handler()
        self.data["источники"][0]["статус_извлечения"] = "извлечено частично"
        self.limitation("Весь источник", "Извлечённые фрагменты ещё не проверены на ошибки. "
                        "Полнота видимого текста требует отдельной визуальной сверки; "
                        "количество фрагментов не равно количеству проверенных областей.",
                        "границы результата")
        return self.data

    def pdf(self):
        import fitz

        with fitz.open(self.path) as doc:
            if doc.needs_pass:
                raise ValueError("PDF защищён паролем; содержимое не извлечено")
            self.data["источники"][0]["страниц"] = len(doc)
            for number, page in enumerate(doc, 1):
                place = f"Страница {number}"
                before = len(self.data["фрагменты"])
                for block in page.get_text("dict", sort=True)["blocks"]:
                    if block.get("type") != 0:
                        continue
                    lines = []
                    for line in block.get("lines", []):
                        lines.append("".join(s["text"] for s in line.get("spans", [])))
                    self.fragment(place + f" → текстовый блок {block.get('number', 0) + 1}",
                                  "\n".join(lines), "текстовый слой PDF",
                                  координаты={"единица": "пункт PDF", "прямоугольник": list(block["bbox"]),
                                              "система": "координаты PyMuPDF без поворота страницы"})
                image_count = 0
                for image_count, info in enumerate(page.get_image_info(), 1):
                    self.limitation(place + f" → изображение {image_count}",
                                    "Текст на изображении не распознан; область в пунктах PDF: " +
                                    str(list(info["bbox"])))
                widgets = list(page.widgets() or [])
                for number_widget, widget in enumerate(widgets, 1):
                    self.fragment(place + f" → поле формы {number_widget} ({widget.field_name})",
                                  widget.field_value, "значение поля PDF", False,
                                  координаты={"единица": "пункт PDF", "прямоугольник": list(widget.rect)})
                for number_annot, annot in enumerate(page.annots() or [], 1):
                    self.fragment(place + f" → аннотация {number_annot}",
                                  annot.info.get("content"), "аннотация PDF", True)
                count = len(self.data["фрагменты"]) - before
                self.coverage(place, извлечено_фрагментов=count, изображений=image_count,
                              поворот_страницы=page.rotation)
                if not count:
                    self.limitation(place, "Доступный текст не найден. Нужны рендер страницы и OCR/визуальное чтение.")
            self.limitation("Все страницы", "Текстовый слой может содержать невидимый, дублированный или устаревший OCR. "
                            "Порядок блоков, таблицы, лигатуры, наложения, текст в кривых и видимость полей формы "
                            "нуждаются в сверке с рендером. OCR PDF этим скриптом не выполняется.")

    def docx(self):
        # XML даёт доступ к тексту ссылок, вложенных таблиц и текстовых полей.
        from docx import Document
        Document(self.path)  # Проверка, что пакет читается как документ Word.
        ns = {"w": W}
        with zipfile.ZipFile(self.path) as z:
            names = z.namelist()
            parts = [n for n in names if n == "word/document.xml" or
                     (n.startswith("word/") and n.endswith(".xml") and
                      (Path(n).name.startswith(("header", "footer")) or
                       Path(n).name in {"footnotes.xml", "endnotes.xml", "comments.xml"}))]
            for part in parts:
                root = ET.fromstring(z.read(part))
                parent = {child: elem for elem in root.iter() for child in elem}
                tables = {elem: i for i, elem in enumerate(root.iter(f"{{{W}}}tbl"), 1)}
                base = self._word_part_label(part)
                paragraphs = list(root.iter(f"{{{W}}}p"))
                before = len(self.data["фрагменты"])
                for index, p in enumerate(paragraphs, 1):
                    ancestors = []
                    cursor = p
                    while cursor in parent:
                        cursor = parent[cursor]
                        ancestors.append(cursor)
                    place = base + f" → абзац {index}"
                    table = next((a for a in ancestors if a.tag == f"{{{W}}}tbl"), None)
                    cell = next((a for a in ancestors if a.tag == f"{{{W}}}tc"), None)
                    row = next((a for a in ancestors if a.tag == f"{{{W}}}tr"), None)
                    if table is not None and cell is not None and row is not None:
                        rows = list(table.findall("w:tr", ns))
                        cells = list(row.findall("w:tc", ns))
                        if row in rows and cell in cells:
                            place = (base + f" → таблица {tables[table]} → строка {rows.index(row) + 1}, "
                                     f"ячейка XML {cells.index(cell) + 1} → абзац {index}")
                    if any(a.tag == f"{{{W}}}txbxContent" for a in ancestors):
                        place += " → текстовое поле"
                    text = []
                    for node in p.iter():
                        nearest = parent.get(node)
                        while nearest is not None and nearest.tag != f"{{{W}}}p":
                            nearest = parent.get(nearest)
                        if nearest is not p:
                            continue  # Вложенное текстовое поле имеет собственный абзац.
                        if node.tag in {f"{{{W}}}t", f"{{{W}}}delText"}:
                            text.append(node.text or "")
                        elif node.tag == f"{{{W}}}tab":
                            text.append("\t")
                        elif node.tag == f"{{{W}}}noBreakHyphen":
                            text.append("\u2011")
                        elif node.tag in {f"{{{W}}}br", f"{{{W}}}cr"}:
                            text.append("\n")
                    revisions = any(a.tag in {f"{{{W}}}del", f"{{{W}}}ins"} for a in ancestors) or any(
                        n.tag in {f"{{{W}}}del", f"{{{W}}}ins", f"{{{W}}}delText"} for n in p.iter())
                    hidden_text = any(n.tag in {f"{{{W}}}vanish", f"{{{W}}}webHidden"}
                                      and n.get(f"{{{W}}}val", "true") not in {"0", "false", "off"}
                                      for n in p.iter())
                    service = Path(part).name == "comments.xml" or revisions or hidden_text
                    style = p.find("w:pPr/w:pStyle", ns)
                    extra = {"часть_пакета": part}
                    if style is not None:
                        extra["стиль_абзаца"] = style.get(f"{{{W}}}val")
                    self.fragment(place, "".join(text), "XML DOCX", service, **extra)
                    if revisions:
                        self.limitation(place, "В абзаце есть записанные исправления. Вставки и удаления могут "
                                        "оказаться вместе; текущую редакцию нужно установить в Word.", "неоднозначная редакция")
                    if hidden_text:
                        self.limitation(place, "Абзац содержит явно скрытые фрагменты. Их видимость нужно "
                                        "установить в Word; весь абзац помечен как служебный.", "скрытый текст")
                    if any(n.tag in {f"{{{W}}}fldSimple", f"{{{W}}}instrText"} for n in p.iter()):
                        self.limitation(place, "Извлечено сохранённое значение поля Word; актуальность поля не проверена.",
                                        "значение не пересчитано")
                tags = {"drawing": "рисунок или схема", "pict": "VML-объект", "altChunk": "внешний фрагмент", "object": "встроенный объект"}
                for tag, label in tags.items():
                    for index, _ in enumerate(root.iter(f"{{{W}}}{tag}"), 1):
                        self.limitation(base + f" → {label} {index}", "Извлечение текста внутри объекта не гарантировано.")
                equations = list(root.iter("{http://schemas.openxmlformats.org/officeDocument/2006/math}oMath"))
                for index, _ in enumerate(equations, 1):
                    self.limitation(base + f" → формула {index}", "Текст математического объекта не извлечён.")
                self.coverage(base, абзацев_в_XML=len(paragraphs),
                              извлечено_фрагментов=len(self.data["фрагменты"]) - before)
            for part in names:
                if part.startswith(("word/charts/", "word/diagrams/", "word/embeddings/")) and not part.endswith("/"):
                    self.limitation(part, "Связанный объект не разобран; текст и подписи требуют отдельной проверки.")
        self.limitation("Весь документ", "Номера страниц без рендера неизвестны. Нумерация списков, скрытый текст, "
                        "условные колонтитулы, разрывы, объединение ячеек и переполнение текста не определены. "
                        "Адрес «ячейка XML» обозначает порядок узлов, а не визуальный номер столбца.")

    @staticmethod
    def _word_part_label(part):
        name = Path(part).name
        if name == "document.xml":
            return "Основной текст"
        if name.startswith("header"):
            return "Верхний колонтитул " + name
        if name.startswith("footer"):
            return "Нижний колонтитул " + name
        return {"footnotes.xml": "Сноски", "endnotes.xml": "Концевые сноски",
                "comments.xml": "Комментарии"}.get(name, name)

    def pptx(self):
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE

        deck = Presentation(self.path)
        self.data["источники"][0]["слайдов"] = len(deck.slides)

        def walk(shapes, base, hidden=False, service=False):
            for shape in shapes:
                place = base + f" → объект {shape.shape_id} «{shape.name}»"
                extra = {"скрытый_слайд": hidden,
                         "координаты": {"единица": "EMU", "лево": shape.left, "верх": shape.top,
                                        "ширина": shape.width, "высота": shape.height}}
                if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                    walk(shape.shapes, place, hidden, service)
                    self.limitation(place, "Координаты дочерних объектов заданы в пространстве группы; "
                                    "для положения на слайде требуется преобразование.", "координаты группы")
                if shape.has_text_frame:
                    for number, paragraph in enumerate(shape.text_frame.paragraphs, 1):
                        self.fragment(place + f" → абзац {number}", paragraph.text,
                                      "текстовый объект PPTX", service or hidden, **extra)
                if shape.has_table:
                    for row_number, row in enumerate(shape.table.rows, 1):
                        for col_number, cell in enumerate(row.cells, 1):
                            if cell.is_spanned:
                                continue
                            self.fragment(place + f" → строка {row_number}, столбец {col_number}",
                                          cell.text, "таблица PPTX", service or hidden, **extra)
                if shape.shape_type in {MSO_SHAPE_TYPE.PICTURE, MSO_SHAPE_TYPE.LINKED_PICTURE}:
                    self.limitation(place, "Текст на изображении не распознан.")
                if shape.has_chart:
                    self.limitation(place, "Диаграмма: заголовок, оси, подписи, легенда и встроенные данные "
                                    "требуют отдельного извлечения и сверки.")
                # SmartArt, OLE, video и иные graphicFrame не равны обычному тексту.
                if not (shape.has_text_frame or shape.has_table or shape.has_chart or
                        shape.shape_type in {MSO_SHAPE_TYPE.GROUP, MSO_SHAPE_TYPE.PICTURE, MSO_SHAPE_TYPE.LINKED_PICTURE}):
                    self.limitation(place, "Объект без доступного текстового блока; возможный текст не извлечён.")

        for number, slide in enumerate(deck.slides, 1):
            place = f"Слайд {number}"
            hidden = slide._element.get("show") in {"0", "false", "off"}
            before = len(self.data["фрагменты"])
            walk(slide.shapes, place, hidden)
            if slide.has_notes_slide:
                notes = slide.notes_slide.notes_text_frame
                if notes is not None:
                    self.fragment(place + " → заметки докладчика", notes.text, "заметки PPTX", True,
                                  скрытый_слайд=hidden)
            self.coverage(place, скрытый_слайд=hidden, извлечено_фрагментов=len(self.data["фрагменты"]) - before)
        # Шаблонные тексты отделены от текста конкретных слайдов.
        for i, master in enumerate(deck.slide_masters, 1):
            walk(master.shapes, f"Образец слайдов {i}", service=True)
            for j, layout in enumerate(master.slide_layouts, 1):
                walk(layout.shapes, f"Образец слайдов {i} → макет {j} «{layout.name}»", service=True)
        with zipfile.ZipFile(self.path) as z:
            for name in z.namelist():
                if name.startswith(("ppt/diagrams/", "ppt/comments/", "ppt/embeddings/")) and not name.endswith("/"):
                    self.limitation(name, "Содержимое части пакета не извлечено; требуется отдельное чтение.")
        self.limitation("Все слайды", "Образцы и макеты извлечены как служебные тексты; их видимость на конкретных "
                        "слайдах не установлена. Анимация, обрезка, слои, текст на фоне и скрытые объекты "
                        "требуют визуальной сверки. Порядок объектов не гарантирует порядок чтения.")

    def xlsx(self):
        from openpyxl import load_workbook
        from openpyxl.cell.cell import MergedCell

        formulas = load_workbook(self.path, data_only=False, read_only=False, keep_links=False)
        cached = load_workbook(self.path, data_only=True, read_only=False, keep_links=False)
        try:
            self.data["источники"][0]["листов"] = len(formulas.worksheets)
            for ws in formulas.worksheets:
                place = f"Лист «{ws.title}»"
                before = len(self.data["фрагменты"])
                hidden = ws.sheet_state != "visible"
                nonempty = 0
                # Все загруженные ячейки без прямоугольного прохода по миллионам пустых строк.
                # _cells — внутренний API openpyxl; проверяется тестом для дальней ячейки.
                for (row, col), cell in sorted(ws._cells.items()):
                    if isinstance(cell, MergedCell):
                        continue
                    cell_place = place + f" → {cell.coordinate}"
                    row_hidden = bool(ws.row_dimensions.get(row) and ws.row_dimensions[row].hidden)
                    col_hidden = any(d.hidden and (d.min or 0) <= col <= (d.max or 0)
                                     for d in ws.column_dimensions.values())
                    extra = {"скрытый_лист": hidden, "скрытая_строка": row_hidden,
                             "скрытый_столбец": col_hidden, "числовой_формат": cell.number_format}
                    service = hidden or row_hidden or col_hidden
                    if cell.value is not None:
                        nonempty += 1
                        if cell.data_type == "f":
                            formula_text = getattr(cell.value, "text", cell.value)
                            value = cached[ws.title][cell.coordinate].value
                            self.fragment(cell_place + " → формула", formula_text, "формула XLSX", True, **extra)
                            if value is not None:
                                self.fragment(cell_place + " → сохранённый результат", self._value(value),
                                              "кэш формулы XLSX", service, **extra)
                            self.limitation(cell_place, "Формула не пересчитана. " +
                                            ("Сохранённый результат отсутствует." if value is None else
                                             "Сохранённый результат может быть устаревшим."), "значение не пересчитано")
                        else:
                            self.fragment(cell_place, self._value(cell.value), "значение ячейки XLSX", service, **extra)
                    if cell.comment is not None:
                        self.fragment(cell_place + " → комментарий", cell.comment.text, "комментарий XLSX", True)
                for area in ws.merged_cells.ranges:
                    self.limitation(place + f" → объединение {area}",
                                    "Значение хранится только в верхней левой ячейке объединения.", "адрес объединения")
                for i, obj in enumerate(ws._images, 1):
                    self.limitation(place + f" → изображение {i}" + self._anchor(obj), "Текст на изображении не распознан.")
                for i, obj in enumerate(ws._charts, 1):
                    self.limitation(place + f" → диаграмма {i}" + self._anchor(obj),
                                    "Заголовки, оси, легенда и подписи диаграммы не извлечены.")
                for label, header_footer in [("обычный верхний колонтитул", ws.oddHeader),
                                              ("обычный нижний колонтитул", ws.oddFooter),
                                              ("чётный верхний колонтитул", ws.evenHeader),
                                              ("чётный нижний колонтитул", ws.evenFooter),
                                              ("первый верхний колонтитул", ws.firstHeader),
                                              ("первый нижний колонтитул", ws.firstFooter)]:
                    for side, side_ru in [("left", "слева"), ("center", "в центре"), ("right", "справа")]:
                        self.fragment(place + f" → {label} → {side_ru}", getattr(header_footer, side).text,
                                      "колонтитул XLSX с кодами полей", True)
                self.coverage(place, непустых_ячеек=nonempty, скрытый_лист=hidden,
                              извлечено_фрагментов=len(self.data["фрагменты"]) - before)
            with zipfile.ZipFile(self.path) as z:
                for name in z.namelist():
                    if name.startswith(("xl/drawings/", "xl/diagrams/", "xl/embeddings/", "xl/chartsheets/", "xl/threadedComments/")) and name.endswith(".xml"):
                        self.limitation(name, "Часть пакета может содержать фигуры, связанный текст или служебные объекты, "
                                        "не представленные в ячейках; требуется отдельная проверка.")
            self.limitation("Вся книга", "Извлечены хранимые значения, а не отображение Excel. Числовые форматы, "
                            "условное форматирование, фильтры, текстовые поля, сводные таблицы и области печати "
                            "требуют визуальной сверки. Внешние ссылки не открываются; формулы не вычисляются. "
                            "Примечания представлены при поддержке openpyxl; современные обсуждения перечислены отдельно.")
        finally:
            formulas.close()
            cached.close()

    @staticmethod
    def _value(value):
        if isinstance(value, (date, datetime, time)):
            return value.isoformat()
        if isinstance(value, bool):
            return "ИСТИНА" if value else "ЛОЖЬ"
        return str(value)

    @staticmethod
    def _anchor(obj):
        anchor = getattr(getattr(obj, "anchor", None), "_from", None)
        if anchor is None:
            return ""
        from openpyxl.utils import get_column_letter
        return f" → якорь {get_column_letter(anchor.col + 1)}{anchor.row + 1}"

    def image(self):
        from PIL import Image, ImageOps

        with Image.open(self.path) as original:
            original.load()
            oriented = ImageOps.exif_transpose(original)
            rgba = oriented.convert("RGBA")
            has_transparency = rgba.getchannel("A").getextrema()[0] < 255
            if has_transparency:
                background = Image.new("RGBA", rgba.size, (255, 255, 255, 255))
                im = Image.alpha_composite(background, rgba).convert("RGB")
                self.data["источники"][0]["фон_для_OCR"] = "белый, прозрачность скомпонована с фоном"
                self.limitation("Изображение целиком", "Для OCR прозрачность скомпонована с белым фоном. "
                                "Белый фон может скрыть светлый текст; требуется визуальная сверка "
                                "на контрастном фоне. Исходное изображение не изменено.", "подготовка к OCR")
            else:
                im = oriented.convert("RGB")
            frame_count = getattr(original, "n_frames", 1)
            self.data["источники"][0]["размер_в_пикселях"] = list(im.size)
            self.coverage("Изображение целиком", кадров=frame_count)
            if frame_count > 1:
                self.limitation("Дополнительные кадры", "Содержимое кадров после первого не извлечено.")
            executable = shutil.which("tesseract")
            if not executable:
                self.limitation("Изображение целиком", "Tesseract не найден. Нужно прочитать изображение визуально или другим OCR.")
                return
            try:
                langs_result = subprocess.run([executable, "--list-langs"], capture_output=True, text=True,
                                              timeout=15, check=True)
                langs = set(langs_result.stdout.splitlines()[1:])
                if "rus" not in langs:
                    self.limitation("Изображение целиком", "Русская языковая модель Tesseract «rus» не установлена. "
                                    "OCR не запускался; требуется визуальное чтение.")
                    return
                language = "rus+eng" if "eng" in langs else "rus"
                with tempfile.TemporaryDirectory(prefix="proverka-ocr-") as tmp:
                    input_path = Path(tmp) / "image.png"
                    im.save(input_path)
                    result = subprocess.run([executable, str(input_path), "stdout", "-l", language, "tsv"],
                                            capture_output=True, text=True, encoding="utf-8", timeout=180, check=True)
                groups: dict[tuple, list] = {}
                for row in csv.DictReader(io.StringIO(result.stdout), delimiter="\t"):
                    if not (row.get("text") or "").strip():
                        continue
                    key = tuple(row[k] for k in ["page_num", "block_num", "par_num", "line_num"])
                    groups.setdefault(key, []).append(row)
                for index, words in enumerate(groups.values(), 1):
                    rects = [(int(w["left"]), int(w["top"]), int(w["width"]), int(w["height"])) for w in words]
                    confidences = [float(w["conf"]) for w in words if float(w["conf"]) >= 0]
                    self.fragment(f"Изображение → строка OCR {index}", " ".join(w["text"] for w in words),
                                  "OCR Tesseract, требует сверки", False,
                                  координаты={"единица": "пиксель", "система": "после применения ориентации EXIF",
                                              "прямоугольник": [min(r[0] for r in rects), min(r[1] for r in rects),
                                                                  max(r[0] + r[2] for r in rects), max(r[1] + r[3] for r in rects)]},
                                  уверенность_OCR=min(confidences) if confidences else None)
                self.limitation("Изображение целиком", "OCR может пропускать текст и искажать буквы, знаки и порядок чтения. "
                                "Число уверенности относится к распознаванию, а не к наличию ошибки автора.")
                if not groups:
                    self.limitation("Изображение целиком", "OCR не нашёл текста. Это не подтверждает отсутствие текста на изображении.")
            except (subprocess.SubprocessError, OSError, ValueError) as exc:
                self.limitation("Изображение целиком", "OCR не завершён: " + str(exc), "ошибка извлечения")


def main(argv=None):
    parser = argparse.ArgumentParser(description="Извлечение текста и адресов без изменения входного файла")
    parser.add_argument("input", help="Путь к PDF, DOCX, PPTX, XLSX, JPEG или PNG")
    parser.add_argument("--output", required=True, help="Путь к JSON UTF-8")
    args = parser.parse_args(argv)
    source, output = Path(args.input).expanduser(), Path(args.output).expanduser()
    if source.resolve() == output.resolve() or (source.exists() and output.exists() and source.samefile(output)):
        parser.error("Выходной JSON не может заменять входной файл")
    extraction = Extraction(source)
    result_code = 0
    try:
        data = extraction.run()
    except Exception as exc:
        data = extraction.data
        data["источники"][0]["статус_извлечения"] = "ошибка извлечения"
        extraction.limitation("Весь источник", f"Не удалось завершить извлечение ({type(exc).__name__}): {exc}",
                              "ошибка извлечения")
        result_code = 1
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(data, ensure_ascii=False, indent=2, allow_nan=False) + "\n", encoding="utf-8")
    print(f"Результат сохранён: {output}", file=sys.stderr)
    return result_code


if __name__ == "__main__":
    sys.exit(main())
